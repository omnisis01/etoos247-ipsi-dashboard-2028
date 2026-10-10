# 2028 전형계획 엑셀(85열·사정모형 단위)을 2027 파서 레이아웃(35열·1행=1전형)의 중간 엑셀로 번역한다
"""
왜 파서를 고치지 않고 번역하나 — 사용자 지시 "기본 틀과 기능은 2027 그대로". build_data.py 는
'전체' 시트·4행부터·35열 고정 레이아웃(r[0..34])을 읽고, 그 뒤에 교정 채널·하네스 11종·app.js 가
걸려 있다. 원천 쪽을 2027 모양으로 맞추면 그 전부가 무수정으로 돌아간다.

입력  ../입결 및 인사이트/2028_수시정시_35개대_의치약한수_260608.xlsx  '수시' 시트
      snap27.json  — 2027 data.js 동결본. sigun·최저·경쟁률·입결·추합·기준 이력을 여기서 상속
출력  ../입결 및 인사이트/2028학년도 수시지원의 모든 것_전형계획기준_v1.xlsx  '전체' 시트

규칙은 전부 2026-10-09 실측에서 나왔다(context-notes.md "2028학년도 판 전환 착수" 참조).
- 1단계/2단계 두 행 → 한 행. 인원은 양쪽에 같은 값이라 한 번만 센다. 전형방법은 col21~28 숫자로 합성.
- 최저 "국, 수, 영, 사, 과 중 2개 영역 합 6이내" → "국,수,영,사,과 2합6" (2027 파서는 (\\d)합(\\d+)를 찾는다).
- 소재지는 524/8,522행만 있어 2027 매칭 행의 sigun 을 행 단위로 상속한다.
- 전년대비는 원천에 없다 → 2027 인원과 4키로 대조해 ▲N/▼N/-/신설 을 쓴다.
  (5키 자격 문구가 2027과 달라 파서의 recompute_prev 는 대부분 실패하므로 여기서 확정한다.)
- 경쟁률 열은 2027/2026/2025, 입결·추합·기준 열은 2026/2025/2024 — 지표별 최신 가용 연도.
"""
import json, os, re, sys, collections
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, '..', '입결 및 인사이트', '2028_수시정시_35개대_의치약한수_260608.xlsx')
OUT = os.path.join(HERE, '..', '입결 및 인사이트', '2028학년도 수시지원의 모든 것_전형계획기준_v1.xlsx')
SNAP = os.path.join(HERE, 'snap27.json')
RENAME = os.path.join(HERE, 'tools', 'plan28', 'rename28.json')      # 2027 전형명 → 2028 전형명 (시행계획 PDF 근거)
SUPP = os.path.join(HERE, 'tools', 'plan28', 'supplement28.json')    # 2028 엑셀이 빠뜨린 행 (시행계획 PDF 모집단위표 근거)

s = lambda v: '' if v is None else str(v).strip()
nz = lambda t: re.sub(r'\s', '', t or '')

# ---------------------------------------------------------------- 매칭용 정규화 (파서 _nz_name 과 같은 규칙)
# 실측(2026-10-09): 2028 '기회균형Ⅰ전형'(유니코드 로마자) vs 2027 '기회균형I전형', 2027 학과명 안의
# 줄바꿈 탓에 괄호 제거가 안 되던 것, '교육기회균형' vs '교육기회균형전형' — 이 셋이 미매칭 2,580행의 대부분이었다.
_ROMAN = {'Ⅰ': 'I', 'Ⅱ': 'II', 'Ⅲ': 'III', 'Ⅳ': 'IV'}
def key(t):
    t = nz(t)
    for a, b in _ROMAN.items(): t = t.replace(a, b)
    return t.replace('·', '').replace('ㆍ', '')
def bare(t):                          # 괄호(미닫힘 포함)와 '전형' 접미를 벗긴 비교 키
    t = re.sub(r'\(.*?\)', '', key(t)); t = re.sub(r'\([^)]*$', '', t)
    return re.sub(r'전형$', '', t)
def paren(t): return ''.join(re.findall(r'\((.*?)\)', key(t)))
_nhap = lambda t: (lambda m: m and m.group(1) + '합' + m.group(2))(re.search(r'(\d)합(\d+)', nz(t)))
# 파서(build_data.py _nz_name)와 같은 3키 정규화 — enroll27.json keys3 악수용
def nz_name(t):
    t = re.sub(r'\s', '', t or '')
    t = t.replace('Ⅰ', 'I').replace('Ⅱ', 'II').replace('Ⅲ', 'III').replace('·', '').replace('ㆍ', '')
    return re.sub(r'전형$', '', t)

# ---------------------------------------------------------------- 대학명: 2028 약칭 → 2027 정식명
# 실측 76교 — 1:1 58 · 태그 없는 이름=태그 없는 2027 이름 8 · 약칭 확장 9 · 강원대(강릉원주) 특례 1.
# 2027 은 캠퍼스를 대학명 꼬리표((글로컬)·(세종)·(천안)·(WISE)·(미래)·(ERICA))와 sigun 으로 구분한다.
_ABBR = [(r'여대$', '여자대학교'), (r'외대', '외국어대학교'), (r'^경상국립대$', '경상대학교'),
         (r'^서울과학기술대$', '서울과학기술대학교'), (r'^한국항공대$', '한국항공대학교')]
_EXPLICIT = {
    '강원대(강릉원주)': '강원대학교(강릉)',   # 7행 전부 소재지 강릉·치의예과 — 요강상 강릉캠(2027 분리 유지)
    '강원대(춘천)': '강원대학교(춘천)',
    '한국외대': '한국외국어대학교',          # 2027 은 단일 대학명 + sigun(서울/용인/송도) — 아래 _HUFS_SIGUN
    '한국외대(글로벌)': '한국외국어대학교',
}
_HUFS_SIGUN = {'한국외대': '서울', '한국외대(글로벌)': '용인'}

def uni27(name, uni_set):
    if name in _EXPLICIT: return _EXPLICIT[name]
    m = re.match(r'^(.*?)(\((.*?)\))?$', name)
    base, tag = m.group(1), (m.group(3) or '')
    for pat, rep in _ABBR:
        base = re.sub(pat, rep, base)
    if not base.endswith('대학교'):
        base = re.sub(r'대$', '대학교', base)
    cand = f'{base}({tag})' if tag else base
    if cand in uni_set: return cand
    # 태그 표기가 조금 달라도(ERICA/에리카) 2027 쪽 꼬리표를 찾는다
    if tag:
        for u in uni_set:
            if u.startswith(base + '(') and nz(tag).lower() in nz(u).lower(): return u
    raise SystemExit(f'[중단] 대학명 매핑 실패: {name!r} → {cand!r} 가 2027 에 없다')

# ---------------------------------------------------------------- 최저: 2028 산문 → 2027 압축 서식
# 2028 반영=Y 3,364행 — "N개 영역 합 M" 3,026 · 12패턴 330 · "각N등급" 8. 전부 여기서 처리한다.
_AREA = r'(?:국|수|영|사|과|탐|한국사|직탐|제2외|한|국어|수학|영어)'
def norm_least(t):
    t = s(t)
    if not t: return ''
    pre = ''
    m = re.match(r'^\[([^\]]*필수\s*응시[^\]]*)\]\s*', t)
    if m: pre = m.group(1); t = t[m.end():]
    t = re.sub(r'\((\d)\s*과목\)', r'(\1)', t)             # 탐(2과목) → 탐(2)
    t = re.sub(r'\s*,\s*', ',', t)
    t = re.sub(r'\s+', ' ', t).strip()
    areas = t.split(' 중 ')[0] if ' 중 ' in t else re.match(r'^([가-힣(),0-9]+)', t).group(1)
    n_listed = len([a for a in areas.split(',') if a])
    # "중 N개 영역 합 M이내" / "영역 합 M이내"(전부) / "중 N개 영역 각 M이내" / "각 M등급"
    # '등급합' 과 '합' 은 같은 뜻 — 잔여 91건이 전부 '등급합' 표기였다(2026-10-09 실측)
    t = re.sub(r' 중 (수학 포함 )?(\d)개 영역 (?:등급)?합 (\d+)\s*이내', r' \1\2합\3', t)
    t = re.sub(r'(,|^|\s)영역 (?:등급)?합 (\d+)\s*이내', lambda m: f'{m.group(1)}{n_listed}합{m.group(2)}', t)
    t = re.sub(r' 중 (수학 포함 )?(\d)개 영역 각 (\d+)\s*(?:이내|등급)', r' \1\2개 각\3', t)
    t = re.sub(r'(,|^|\s)영역 각 (\d+)\s*이내', lambda m: f'{m.group(1)}{n_listed}개 각{m.group(2)}', t)
    t = re.sub(r'(\d)\s*등급 (\d)개 이상', r'\1등급 \2개', t)
    t = re.sub(r'(한국사|영어)\s*(\d+)\s*이내', r'\1\2', t)
    t = re.sub(r'\s*이내', '', t)
    t = re.sub(r'\s*\(', ' (', t).strip()
    if pre: t = f'{t} [{nz(pre).replace("수능","")}]'
    return t

# ---------------------------------------------------------------- 전형방법 합성 (2027 어휘)
_COMP = [(21, '교과'), (22, '서류'), (23, '면접'), (24, '논술'), (25, '실기'), (26, '서류'), (27, '1단계성적'), (28, '기타')]
def comp_text(r):
    parts = []
    for ci, name in _COMP:
        v = s(r[ci])
        if not v or v in ('0', '0.0'): continue
        v = v[:-2] if v.endswith('.0') else v
        if parts and parts[-1][0] == name:            # 학생부(정성)와 서류가 둘 다 '서류'면 합친다
            parts[-1] = (name, str(int(float(parts[-1][1])) + int(float(v)))); continue
        parts.append((name, v))
    return '+'.join(f'{n}{v}' for n, v in parts) or s(r[20])

def method_text(stages):
    if len(stages) == 1:
        return comp_text(stages[0])
    st1 = next((r for r in stages if s(r[16]) == '1단계'), None)
    st2 = next((r for r in stages if s(r[16]) == '2단계'), None)
    if not (st1 and st2): return comp_text(stages[0])
    ratio = s(st1[17]); mult = ''
    if ratio.replace('.', '').isdigit() and float(ratio) > 100:
        mult = f'({float(ratio) / 100:g}배수)'
    return f'1단계: {comp_text(st1)}{mult} → 2단계: {comp_text(st2)}'

# ---------------------------------------------------------------- 학년별 비율 (col33~39) · 서류 약어
_GR = [(33, '1학년'), (34, '2학년'), (35, '3학년'), (36, '1:2학년'), (37, '2:3학년'), (38, '1:2:3학년'), (39, '1:3학년')]
def grade_ratio(r):
    out = [f'{n} {s(r[i]).rstrip("0").rstrip(".")}' for i, n in _GR if s(r[i]) not in ('', '0', '0.0')]
    return ' · '.join(out)

_DOC = [('학교생활기록부', '학'), ('지원자격증빙서류', '증'), ('학교장 추천 명단', '추'), ('추천인확인서', '추'),
        ('자기소개서', '자'), ('활동보고서', '활'), ('학업이수계획서', '자')]
def docs_abbr(t):
    out = []
    for w in [x.strip() for x in s(t).split(',') if x.strip()]:
        ab = next((a for k, a in _DOC if k in w), None)
        out.append(ab or w)
    return ','.join(dict.fromkeys(out))

# ---------------------------------------------------------------- 메인
def main():
    snap = json.load(open(SNAP, encoding='utf-8'))
    R27, LOC27 = snap['rows'], snap['uniLoc']
    # 전형명 개명(레이어 C) — 2027 쪽을 2028 이름으로 바꿔 색인한다. 안 그러면 개명된 전형의 전 행이 T7 공란이 되어
    # 파서가 '전형 변경'으로 덮는다(실측 2026-10-10: 서울대 지역균형선발전형→지역균형전형 60행).
    _ren = {u: m for u, m in (json.load(open(RENAME, encoding='utf-8')) if os.path.exists(RENAME) else {}).items() if not u.startswith('_')}
    ren = {u: m.get('jhname', {}) for u, m in _ren.items()}       # 전형명 개명
    ren_d = {u: m.get('dept', {}) for u, m in _ren.items()}       # 모집단위 개명(2027 → 2028 엑셀 표기)
    ren_used = set()
    uni_set = set(LOC27)
    SIGUN2REG = {}
    for _u, _locs in LOC27.items():
        for _l in _locs: SIGUN2REG.setdefault(_l['sigun'], _l['region'])
    # 2027 행 색인 — 정확 4키 · 3키 · 괄호 제거 3키(권역·세부전공 꼬리표 무시)
    by4, by3, byBB, deptsU = {}, collections.defaultdict(list), collections.defaultdict(list), collections.defaultdict(set)
    for k, v in R27.items():
        u, d, jt, jn = k.split('|')
        if d in ren_d.get(u, {}): ren_used.add((u, d)); d = ren_d[u][d]
        if jn in ren.get(u, {}):
            ren_used.add((u, jn)); jn = ren[u][jn]
            if (u, key(d), jt, key(jn)) in by4: raise SystemExit(f'[중단] rename28 충돌: {u} {d} {jt} {jn} 가 이미 2027 에 있다')
        by4[(u, key(d), jt, key(jn))] = v
        by3[(u, key(d), key(jn))].append((jt, v))
        byBB[(u, bare(d), bare(jn), '(외)' in jn)].append((jn, v))
        deptsU[u].add(bare(d))

    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    raw = [r for r in wb['수시'].iter_rows(min_row=3, values_only=True) if s(r[6])]
    groups = collections.OrderedDict()
    for r in raw:
        groups.setdefault((s(r[6]), s(r[9]), s(r[12]), s(r[13])), []).append(r)
    # ---- 보완 행(레이어 A): 2028 엑셀이 빠뜨린 모집단위를 시행계획 PDF 모집단위표 근거로 넣는다(tools/plan28/supplement28.json).
    #      같은 대학·전형의 형제 행(전 단계)을 복사해 공통정보(자격·서류·단계 배수)를 받고, 학과·인원·(명시 시) 전형방법·최저만 바꾼다.
    #      실측(2026-10-10): 서울대 음대·미대·체육교육과 15단위 121명이 2027 엔 있었는데 6/8 엑셀에서 통째로 빠졌다.
    supp = json.load(open(SUPP, encoding='utf-8'))['rows'] if os.path.exists(SUPP) else []
    SUPP_METHOD, SUPP_LEAST, SUPP_NOTE = {}, {}, {}
    for x in supp:
        sib = next((st for (u28, jn28, _, _), st in groups.items() if u28 == x['uni'] and jn28 == x['jhname']), None)
        if sib is None: raise SystemExit(f"[중단] supplement28: 형제 행 없음 {x['uni']} {x['jhname']}")
        gk = (x['uni'], x['jhname'], x['dept'], None)
        if gk in groups: raise SystemExit(f"[중단] supplement28: 원천에 이미 있음 {gk} — 엑셀이 갱신됐다면 제거할 것")
        stages = []
        for r0 in sib:
            r = list(r0); r[12], r[13], r[15] = x['dept'], None, x['enroll']
            if 'least' in x: r[64], r[69] = ('N', None) if x['least'] == '없음' else ('Y', x['least'])
            stages.append(tuple(r))
        groups[gk] = stages
        if x.get('method'): SUPP_METHOD[gk] = x['method']
        if 'least' in x: SUPP_LEAST[gk] = x['least']
        SUPP_NOTE[gk] = f"모집인원은 2028 시행계획 {x['src']} 기준(6월 엑셀 누락 보완)"
    # 학과 해석 — 2027 학과명으로 환원한다. 직접 일치 > '학부-학과' 하이픈 부분 일치(2028 이 상위 학부를
    # 접두로 붙인 경우: '경영학부-경영학' ↔ 2027 '경영학부', '건축학부-건축공학' ↔ '건축공학전공').
    # 실측(2026-10-09): 신설 671행 중 상당수가 이 꼴이었다 — 신설은 화면에서 강한 신호라 가짜를 두면 안 된다.
    def resolve_dept(uni, dept28):
        b = bare(dept28)
        if b in deptsU[uni]: return b
        if '-' not in dept28: return None
        hits = set()
        for part in dept28.split('-'):
            pb = bare(part)
            for cand in (pb, pb + '전공', pb + '과', pb.rstrip('과') + '전공'):
                if cand in deptsU[uni]: hits.add(cand)
        return hits.pop() if len(hits) == 1 else None     # 부분이 서로 다른 2027 학과를 가리키면 포기
    # 2028 쪽 (대학, 해석된 학과, 괄호제거 전형명, 정원외) 그룹 크기 — 1:N 합산·하이픈 환원은 2028 이 1행일 때만.
    # 2028 여러 행이 같은 2027 행으로 가면 같은 인원을 중복 비교하게 되므로 금지한다.
    # ⚠️ g28d 는 2027 학과 하나로 환원되는 2028 **학과명의 집합**이다. 전형별 그룹 수로 세면 전형이 2개 이상인
    #    하이픈 학과가 모두 '분리'가 된다(실측 2026-10-10: 서울대 물리천문학부-물리학 지균·일반·기균 3그룹 → 가짜 분리).
    g28, g28d = collections.Counter(), collections.defaultdict(set)
    for (u28, jn28, dept28, sub28), st in groups.items():
        u = uni27(u28, uni_set); jn = jn28 + ('(외)' if s(st[0][7]) == '정원외' else '')
        rd = resolve_dept(u, dept28)
        g28[(u, rd, bare(jn), '(외)' in jn)] += 1
        if rd is not None and rd != bare(dept28): g28d[(u, rd)].add(dept28)   # 하이픈 환원으로 같은 2027 학과에 모이는 2028 학과들

    out, log, matched_k3 = [], collections.Counter(), set()
    if supp: log['보완 행(시행계획 PDF 모집단위표)'] = len(supp)
    if ren_used: log['전형명 개명 적용(rename28)'] = len(ren_used)
    _ren_miss = [(u, o) for mm in (ren, ren_d) for u, m in mm.items() for o in m if (u, o) not in ren_used]
    if _ren_miss: raise SystemExit(f'[중단] rename28 미적용 {_ren_miss} — 2027 스냅에 그 전형명이 없다. 오타거나 이미 반영됐으면 제거할 것')
    unmatched_least = []
    for (u28, jn28, dept28, sub28), stages in groups.items():
        r = stages[0]
        gk = (u28, jn28, dept28, sub28)
        uni = uni27(u28, uni_set)
        jtype = s(r[8])
        if jtype == '실기': jtype = '특기자' if '특기자' in jn28 else '실기/실적'
        jhname = jn28 + ('(외)' if s(r[7]) == '정원외' else '')
        dept = dept28 + (f'({sub28})' if sub28 and sub28 not in dept28 else '')
        enrolls = {s(x[15]) for x in stages}
        if len(enrolls) > 1: log['인원 불일치(최대값 사용)'] += 1
        enroll = max(int(float(e)) for e in enrolls if e)
        if len(stages) == 3: log['1단계+2단계+일괄합산(일괄합산 무시)'] += 1

        # 2027 매칭 — 단계적. 학과는 2028 원형(dept28, 세부전공 미부착)을 2027 학과로 환원해 찾는다.
        ext = '(외)' in jhname
        rd = resolve_dept(uni, dept28)
        hyph = rd is not None and rd != bare(dept28)                 # 하이픈 환원으로 찾은 학과
        split = hyph and len(g28d[(uni, rd)]) > 1                     # 2028 학과 둘 이상이 같은 2027 학과로 → 분리
        m27, tier = None, None
        if (x := by4.get((uni, key(dept28), jtype, key(jhname)))): m27, tier = x, 'T1 정확'
        elif len(c := by3.get((uni, key(dept28), key(jhname)), [])) == 1: m27, tier = c[0][1], 'T2 전형유형 상이'
        elif rd is None: tier = 'T8 학과 없음→신설'
        elif split: tier = 'T9 하이픈 환원 N:1→분리'                   # 같은 인원을 중복 비교하지 않고 '분리'로 표기
        else:
            cand = byBB.get((uni, rd, bare(jhname), ext), [])
            if len(cand) == 1 and g28[(uni, rd, bare(jhname), ext)] > 1:
                # 2027 전형 하나를 2028 이 권역 등으로 여러 행으로 쪼갰다 — 각 행을 2027 총원과 비교하면
                # 가짜 증감(▼5·▲1·▼3…)이 생기고 경쟁률이 N행에 복사된다(실측: 순천향대·원광대 의예과 6행씩,
                # verify_data 블록 오염 래칫이 잡았다). 하이픈 학과(T9)와 같은 뜻이므로 '분리'로 쓴다.
                tier = 'T10 전형 1:N 분할→분리'
            elif len(cand) == 1: m27, tier = cand[0][1], ('T3b 하이픈 환원 1:1' if hyph else 'T3 괄호제거 1:1')
            elif len(cand) > 1 and g28[(uni, rd, bare(jhname), ext)] == 1:
                # 2027 이 권역별로 쪼갠 것을 2028 이 한 행으로 모았다 — 인원은 합산, 이력은 최대 인원 행
                best = max(cand, key=lambda c: c[1].get('enroll') or 0)[1]
                m27 = dict(best, enroll=sum(c[1].get('enroll') or 0 for c in cand)); tier = 'T4 1:N 합산'
            elif len(cand) > 1:
                same = [c for c in cand if paren(jhname) and paren(jhname) in paren(c[0])]
                if len(same) == 1: m27, tier = same[0][1], 'T5 괄호내용 짝'
                else: tier = 'T6 N:N 짝 실패→공란'
            else: tier = 'T7 학과만 존재→공란(파서가 전형 변경 판정)'
        log[tier] += 1
        if m27: matched_k3.add('|'.join((uni, nz_name(dept), nz_name(jhname))))

        # 전년대비 — 어댑터가 확정한다. 공란은 파서가 keys3(enroll27.json)로 '전형 변경'을 판정하고,
        # '신설'은 2027 에 그 학과 자체가 없을 때만, '분리'는 2027 한 학과가 2028 여러 학과로 갈린 때만 쓴다.
        if m27 and m27.get('enroll') is not None:
            d = enroll - int(m27['enroll'])
            prev = '-' if d == 0 else (f'▲{d}' if d > 0 else f'▼{-d}')
        elif tier.startswith('T8'): prev = '신설'
        elif tier.startswith(('T9', 'T10')): prev = '분리'
        else: prev = ''

        # 최저 — Y/N 과 산문 정규화, 2027 과 다르면 변경사항 합성(parse_choejeo_change 가 읽는 서식)
        least = norm_least(r[69]) if s(r[64]) == 'Y' else '없음'
        if gk in SUPP_LEAST: least = SUPP_LEAST[gk]                   # 보완 행은 PDF 조항을 그대로 쓴다(정규화 생략)
        if s(r[64]) == 'Y' and not least: least = '없음'; log['최저 Y인데 내용 없음'] += 1
        if s(r[64]) == 'Y' and least != '없음' and not re.search(r'\d합\d|각\d|등급 \d개', least):
            unmatched_least.append(s(r[69]))
        change = ''
        if m27:
            l27 = s(m27.get('choejeo')) or '없음'
            # 2027 '국,수,영,탐(1) 2합6' 와 2028 정규화 '국,수,영,사,과 2합6' 는 같은 최저다. 문구로 비교하면
            # 553행이 가짜 '변경'이 된다(2026-10-09 실측). N합M 이 다를 때만 변화로 본다. 하위 조건(수(기미) 지정·
            # 탐구 과목 수)은 두 표기법 사이에서 신뢰성 있게 비교할 수 없어 가짜 신호 대신 침묵을 택한다.
            h27, h28 = _nhap(l27), _nhap(least)
            if l27 == '없음' and least != '없음': change = '수능최저 신설'
            elif l27 != '없음' and least == '없음': change = '수능최저 폐지'
            elif h27 and h28 and h27 != h28: change = f'최저: {l27} → {least}'
            elif not h27 and not h28 and nz(l27) != nz(least): change = f'최저: {l27} → {least}'   # 각N 형 등
        # 개명된 전형은 변경사항에 옛 이름을 남긴다(parse_choejeo_change 는 '최저·합·등급' 없는 구간을 무시한다)
        _old = next((o for o, n in ren.get(uni, {}).items() if n == jhname), None)
        if _old: change = (change + ' / ' if change else '') + f'전형명 변경: {_old} → {jhname}'
        _oldd = next((o for o, n in ren_d.get(uni, {}).items() if n == dept28), None)
        if _oldd: change = (change + ' / ' if change else '') + f'모집단위명 변경: {_oldd} → {dept28}'

        # 소재지 — 2028 소재지 > 한국외대 특례 > 2027 매칭 행 > 2027 대학 최빈
        sigun = s(r[14]) or _HUFS_SIGUN.get(u28) or (m27 and m27.get('sigun')) or LOC27[uni][0]['sigun']
        # 광역은 기초에서 역산한다 — 2028 지역(col3)은 대학 단위라 소재지가 분캠이면 '서울|용인' 같은 짝이 생긴다(실측 300행).
        # 2027 그 대학의 (광역,기초) → 2027 전체의 기초→광역 → 2028 col3 순.
        hit = next((x for x in LOC27[uni] if x['sigun'] == sigun), None)
        region = hit['region'] if hit else SIGUN2REG.get(sigun) or s(r[3])

        h = m27 or {}
        c = h.get('c', [None] * 3); g = h.get('g', [None] * 3); v = h.get('v', [None] * 3)
        ch = h.get('chung', [''] * 3); sd = h.get('std', [''] * 3)
        out.append([
            region, sigun, uni, s(r[10]), dept, jtype, jhname, s(r[31]), enroll, prev, change,
            least, SUPP_METHOD.get(gk) or method_text(stages), docs_abbr(r[79]), '', grade_ratio(r), s(r[49]), s(r[61]),
            c[0], c[1], c[2],
            sd[0], g[0], v[0], ch[0], SUPP_NOTE.get(gk, ''),
            sd[1], g[1], v[1], ch[1],
            sd[2], g[2], v[2], ch[2],
            '',
        ])

    # ---- 자체 검증: 틀린 중간 산출물로 파서를 돌리면 사고가 조용히 다음 단계로 번진다
    fails = []
    if len(out) != len(groups): fails.append(f'행수 {len(out)} ≠ 고유키 {len(groups)}')
    src_sum = sum(int(float(s(r[15]))) for (k, st) in groups.items() for r in st[:1] if s(r[15]))
    if sum(x[8] for x in out) != src_sum: fails.append(f'인원 합 {sum(x[8] for x in out)} ≠ 원천 {src_sum}')
    y_src = sum(1 for st in groups.values() if s(st[0][64]) == 'Y')
    y_out = sum(1 for x in out if x[11] != '없음')
    if y_src != y_out: fails.append(f'최저 있음 {y_out} ≠ 원천 Y {y_src}')
    if unmatched_least: fails.append(f'최저 정규화 잔여 {len(unmatched_least)}건 — 예: {unmatched_least[:3]}')
    # 대학명 매핑 실패는 uni27() 이 즉시 중단시킨다. 2028 이름 수 > 2027 이름 수인 것은
    # 한국외대·한국외대(글로벌) → 한국외국어대학교(sigun 으로 구분) 같은 **의도된 병합**이다.
    merged = len({k[0] for k in groups}) - len({x[2] for x in out})
    if merged: log[f'대학명 병합(2028 {len({k[0] for k in groups})}→2027 {len({x[2] for x in out})}교)'] = merged

    # enroll27.json — 파서가 읽는 작년 스냅샷. enroll27(5키)은 2027 인원, keys3 는 2027 3키에 **어댑터가 매칭한 2028 행의
    # 3키를 더한 것**이다. 파서의 is_changed_track 은 3키가 keys3 에 없으면 '전형 변경'으로 덮는데, 2027 전형명엔
    # '(경기도 의정부권)' 같은 꼬리표가 있어 어댑터가 괄호 제거로 맞춘 행도 전부 덮였다(실측 525행). 이 악수로 막는다.
    e27 = {'|'.join([u, d, jt, jn, (v.get('jagyeok') or '')]): v['enroll'] for k, v in R27.items()
           for u, d, jt, jn in [k.split('|')] if v.get('enroll') is not None}
    k3 = {'|'.join((u, nz_name(d), nz_name(jn))) for k in R27 for u, d, jt, jn in [k.split('|')]} | matched_k3
    json.dump({'meta': {'source': snap['meta']['source'], 'frozen': snap['meta']['frozen'],
        'field': 'enroll27: 2027 data.js 모집인원(5키) / keys3: 2027 3키(_nz_name) ∪ adapt_2028 이 매칭한 2028 행의 3키 — '
                 'build_data.py is_changed_track 과의 악수. verify_insights from 축도 읽는다'},
        'enroll27': e27, 'keys3': sorted(k3)}, open(os.path.join(HERE, 'enroll27.json'), 'w', encoding='utf-8'), ensure_ascii=False)
    print(f'  enroll27.json: 5키 {len(e27)} · keys3 {len(k3)} (2027 {len(k3) - len(matched_k3 - set())} + 매칭 2028 {len(matched_k3)})')

    wbo = openpyxl.Workbook(); ws = wbo.active; ws.title = '전체'
    ws.append(['2028학년도 수시 (전형계획 기준 — adapt_2028.py 가 2027 레이아웃으로 번역)'] + [''] * 34)
    ws.append([''] * 35)
    ws.append(['광역', '기초', '대학교', '계열', '모집단위명', '전형유형', '전형명', '지원자격', '모집 인원', '전년 대비',
               '전년대비 변경사항', '최저학력기준', '전형방법', '필요 서류', '복수 지원', '학년별반영비율', '반영과목', '진로선택과목',
               '2027학년도 경쟁률', '2026학년도 경쟁률', '2025학년도 경쟁률',
               '2026학년도 기준', '2026학년도 입결(등급)', '2026학년도 입결(환산점수)', '2026 충원', '지원시 유의사항',
               '2025학년도 기준2', '2025학년도 입결(등급)', '2025학년도 입결(환산점수)', '2025 충원',
               '2024학년도 기준2', '2024학년도 입결(등급)', '2024학년도 입결(환산점수)', '2024 충원', '대학별고사 실시일'])
    for row in out: ws.append(row)
    wbo.save(OUT)

    print(f'[adapt_2028] {len(raw)}행 → {len(out)}행 · {len({x[2] for x in out})}교 · 인원 합 {sum(x[8] for x in out):,}')
    for k, n in log.most_common(): print(f'  {k}: {n}')
    print(f'  전년대비 분포: {dict(collections.Counter("공란(전형 변경 후보)" if x[9]=="" else x[9] if x[9] in ("신설","-") else x[9][0] for x in out))}')
    print(f'  최저 변경사항 합성: {sum(1 for x in out if x[10])}건')
    print(f'  → {os.path.basename(OUT)}')
    if fails:
        print('✗ 자체 검증 실패 — 출력은 썼지만 파서에 넣지 마라'); [print('   ', f) for f in fails]; sys.exit(1)
    print('OK  자체 검증 통과(행수·인원 합·최저 Y/N·정규화 잔여·대학 수)')

if __name__ == '__main__':
    main()
